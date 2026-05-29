"""
csv_merge.py - Merge CSV files in a folder into one CSV file.

Examples:
  python csv_merge.py C:/work/downloads --output merged.csv
  python csv_merge.py --downloads --sort mtime --output ../../csv/download_merged.csv

Default behavior:
  - Reads *.csv files from the target folder.
  - Sorts files by file name.
  - Writes the header row from the first CSV only.
  - Skips the header row from the second and later CSV files.
  - Keeps every other row exactly as CSV rows.
"""

from __future__ import annotations

import argparse
import csv
import shutil
import sys
import time
from pathlib import Path


def get_downloads_dir() -> Path:
    return Path.home() / "Downloads"


def collect_csv_files(
    folder: Path,
    pattern: str,
    sort_mode: str,
    output_path: Path | None,
    since_minutes: float | None,
) -> list[Path]:
    files = [p for p in folder.glob(pattern) if p.is_file() and p.suffix.lower() == ".csv"]

    if since_minutes is not None:
        threshold = time.time() - (since_minutes * 60)
        files = [p for p in files if p.stat().st_mtime >= threshold]

    if output_path is not None:
        try:
            output_resolved = output_path.resolve()
            files = [p for p in files if p.resolve() != output_resolved]
        except OSError:
            pass

    if sort_mode == "name":
        files.sort(key=lambda p: p.name.lower())
    elif sort_mode == "mtime":
        files.sort(key=lambda p: (p.stat().st_mtime, p.name.lower()))
    else:
        raise ValueError(f"Unsupported sort mode: {sort_mode}")

    return files


def sniff_dialect(file_path: Path, encoding: str) -> csv.Dialect:
    with file_path.open("r", newline="", encoding=encoding) as f:
        sample = f.read(4096)
    try:
        return csv.Sniffer().sniff(sample)
    except csv.Error:
        return csv.excel


def merge_csv_files(
    files: list[Path],
    output_path: Path,
    encoding: str,
    output_encoding: str,
    keep_all_headers: bool,
    no_header_output: bool,
    add_source_column: bool,
) -> tuple[int, int]:
    total_rows = 0
    merged_files = 0
    first_header: list[str] | None = None
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with output_path.open("w", newline="", encoding=output_encoding) as out:
        writer: csv.writer | None = None

        for file_path in files:
            dialect = sniff_dialect(file_path, encoding)
            with file_path.open("r", newline="", encoding=encoding) as f:
                reader = csv.reader(f, dialect)
                try:
                    header = next(reader)
                except StopIteration:
                    print(f"[WARN] Empty CSV skipped: {file_path}", file=sys.stderr)
                    continue

                if writer is None:
                    writer = csv.writer(out, dialect=dialect)
                    first_header = list(header)
                    if not no_header_output:
                        output_header = list(header)
                        if add_source_column:
                            output_header.insert(0, "source_file")
                        writer.writerow(output_header)
                elif keep_all_headers:
                    row = list(header)
                    if add_source_column:
                        row.insert(0, file_path.name)
                    writer.writerow(row)
                    total_rows += 1
                elif first_header is not None and header != first_header:
                    print(
                        f"[WARN] Header differs from first CSV and was skipped: {file_path.name}",
                        file=sys.stderr,
                    )

                for row in reader:
                    if add_source_column:
                        row = [file_path.name] + row
                    writer.writerow(row)
                    total_rows += 1

                merged_files += 1

    return merged_files, total_rows


def build_merged_rows(
    files: list[Path],
    encoding: str,
    keep_all_headers: bool,
    no_header_output: bool,
    add_source_column: bool,
) -> tuple[csv.Dialect, list[list[str]], int]:
    output_rows: list[list[str]] = []
    total_data_rows = 0
    first_header: list[str] | None = None
    output_dialect: csv.Dialect | None = None

    for file_path in files:
        dialect = sniff_dialect(file_path, encoding)
        if output_dialect is None:
            output_dialect = dialect

        with file_path.open("r", newline="", encoding=encoding) as f:
            reader = csv.reader(f, dialect)
            try:
                header = next(reader)
            except StopIteration:
                print(f"[WARN] Empty CSV skipped: {file_path}", file=sys.stderr)
                continue

            if first_header is None:
                first_header = list(header)
                if not no_header_output:
                    output_header = list(header)
                    if add_source_column:
                        output_header.insert(0, "source_file")
                    output_rows.append(output_header)
            elif keep_all_headers:
                row = list(header)
                if add_source_column:
                    row.insert(0, file_path.name)
                output_rows.append(row)
                total_data_rows += 1
            elif header != first_header:
                print(
                    f"[WARN] Header differs from first CSV and was skipped: {file_path.name}",
                    file=sys.stderr,
                )

            for row in reader:
                if add_source_column:
                    row = [file_path.name] + row
                output_rows.append(row)
                total_data_rows += 1

    if output_dialect is None:
        output_dialect = csv.excel

    return output_dialect, output_rows, total_data_rows


def write_rows_to_csv(rows: list[list[str]], output_path: Path, dialect: csv.Dialect, encoding: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding=encoding) as out:
        writer = csv.writer(out, dialect=dialect)
        writer.writerows(rows)


def paste_rows_to_csv(
    rows: list[list[str]],
    target_path: Path,
    start_row: int,
    mode: str,
    encoding: str,
    output_encoding: str,
    backup: bool,
) -> int:
    if start_row < 1:
        raise ValueError("--paste-start-row must be 1 or greater.")
    if mode not in {"replace", "insert"}:
        raise ValueError("--paste-mode must be replace or insert.")
    if not target_path.is_file():
        raise FileNotFoundError(f"Paste target file not found: {target_path}")

    dialect = sniff_dialect(target_path, encoding)
    with target_path.open("r", newline="", encoding=encoding) as f:
        target_rows = list(csv.reader(f, dialect))

    index = start_row - 1
    if len(target_rows) < index:
        target_rows.extend([[] for _ in range(index - len(target_rows))])

    if mode == "insert":
        new_rows = target_rows[:index] + rows + target_rows[index:]
    else:
        new_rows = target_rows[:index] + rows + target_rows[index + len(rows):]

    if backup:
        backup_path = target_path.with_suffix(target_path.suffix + ".bak")
        shutil.copy2(target_path, backup_path)

    write_rows_to_csv(new_rows, target_path, dialect, output_encoding)
    return len(rows)


def parse_excel_column(value: str) -> int:
    value = value.strip()
    if value.isdigit():
        col = int(value)
        if col < 1:
            raise ValueError("--paste-start-col must be 1 or greater.")
        return col

    col = 0
    for char in value.upper():
        if not ("A" <= char <= "Z"):
            raise ValueError("--paste-start-col must be a column number or letter, such as 1 or A.")
        col = col * 26 + (ord(char) - ord("A") + 1)
    return col


def paste_rows_to_excel(
    rows: list[list[str]],
    target_path: Path,
    sheet_name: str | None,
    start_row: int,
    start_col: int,
    mode: str,
) -> int:
    if start_row < 1:
        raise ValueError("--paste-start-row must be 1 or greater.")
    if start_col < 1:
        raise ValueError("--paste-start-col must be 1 or greater.")
    if mode not in {"replace", "insert"}:
        raise ValueError("--paste-mode must be replace or insert.")
    if not target_path.is_file():
        raise FileNotFoundError(f"Excel file not found: {target_path}")

    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel paste. Install it with: pip install openpyxl") from exc

    workbook = openpyxl.load_workbook(target_path)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active

    if mode == "insert" and rows:
        worksheet.insert_rows(start_row, amount=len(rows))

    for row_offset, row_values in enumerate(rows):
        for col_offset, value in enumerate(row_values):
            worksheet.cell(
                row=start_row + row_offset,
                column=start_col + col_offset,
                value=value,
            )

    workbook.save(target_path)
    workbook.close()
    return len(rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge CSV files in a folder into one CSV file.",
    )
    parser.add_argument(
        "folder",
        nargs="?",
        help="Folder that contains CSV files. Omit with --downloads to use the Downloads folder.",
    )
    parser.add_argument(
        "--downloads",
        action="store_true",
        help="Use the current user's Downloads folder as the input folder.",
    )
    parser.add_argument(
        "--pattern",
        default="*.csv",
        help="File name pattern to merge. Default: *.csv",
    )
    parser.add_argument(
        "--output",
        "-o",
        default="merged.csv",
        help="Output CSV path. Default: merged.csv",
    )
    parser.add_argument(
        "--sort",
        choices=("name", "mtime"),
        default="name",
        help="Merge order: name or mtime. Default: name",
    )
    parser.add_argument(
        "--since-minutes",
        type=float,
        default=None,
        help="Merge only CSV files modified within the last N minutes. Example: --since-minutes 5",
    )
    parser.add_argument(
        "--encoding",
        default="utf-8-sig",
        help="Input CSV encoding. Default: utf-8-sig",
    )
    parser.add_argument(
        "--output-encoding",
        default=None,
        help="Output CSV encoding. Default: same as --encoding",
    )
    parser.add_argument(
        "--keep-all-headers",
        action="store_true",
        help="Keep header rows from every CSV. Default: keep only the first header.",
    )
    parser.add_argument(
        "--no-header-output",
        action="store_true",
        help="Do not write the first row/header to the output CSV.",
    )
    parser.add_argument(
        "--add-source-column",
        action="store_true",
        help="Add source_file as the first column.",
    )
    parser.add_argument(
        "--paste-to",
        default=None,
        help="Paste merged rows into this CSV file.",
    )
    parser.add_argument(
        "--paste-to-excel",
        default=None,
        help="Paste merged rows as values into this Excel .xlsx file.",
    )
    parser.add_argument(
        "--excel-sheet",
        default=None,
        help="Excel sheet name for --paste-to-excel. Default: active sheet.",
    )
    parser.add_argument(
        "--paste-start-row",
        type=int,
        default=None,
        help="1-based row number where merged rows are pasted. Required with --paste-to or --paste-to-excel.",
    )
    parser.add_argument(
        "--paste-start-col",
        default="A",
        help="Excel start column for --paste-to-excel. Column letter or number. Default: A",
    )
    parser.add_argument(
        "--paste-mode",
        choices=("replace", "insert"),
        default="replace",
        help="replace: overwrite rows from start row. insert: insert rows before start row. Default: replace",
    )
    parser.add_argument(
        "--no-backup",
        action="store_true",
        help="Do not create a .bak backup when using --paste-to.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    if args.downloads:
        folder = get_downloads_dir()
    elif args.folder:
        folder = Path(args.folder)
    else:
        print("[ERROR] Specify a folder or use --downloads.", file=sys.stderr)
        sys.exit(1)

    if not folder.is_dir():
        print(f"[ERROR] Folder not found: {folder}", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = Path.cwd() / output_path

    output_encoding = args.output_encoding or args.encoding
    if args.since_minutes is not None and args.since_minutes <= 0:
        print("[ERROR] --since-minutes must be greater than 0.", file=sys.stderr)
        sys.exit(1)
    if (args.paste_to or args.paste_to_excel) and args.paste_start_row is None:
        print("[ERROR] --paste-start-row is required with --paste-to or --paste-to-excel.", file=sys.stderr)
        sys.exit(1)
    if args.paste_to and args.paste_to_excel:
        print("[ERROR] Use either --paste-to or --paste-to-excel, not both.", file=sys.stderr)
        sys.exit(1)

    files = collect_csv_files(folder, args.pattern, args.sort, output_path, args.since_minutes)

    if not files:
        condition = f", modified within {args.since_minutes} minute(s)" if args.since_minutes else ""
        print(f"No CSV files found: {folder} ({args.pattern}{condition})")
        sys.exit(0)

    print(f"Input folder : {folder}")
    print(f"Output file  : {output_path}")
    print(f"Sort order   : {args.sort}")
    if args.since_minutes:
        print(f"Modified in  : last {args.since_minutes:g} minute(s)")
    print(f"CSV files    : {len(files)}")
    for file_path in files:
        print(f"  - {file_path.name}")
    print()

    dialect, merged_rows, total_rows = build_merged_rows(
        files=files,
        encoding=args.encoding,
        keep_all_headers=args.keep_all_headers,
        no_header_output=args.no_header_output,
        add_source_column=args.add_source_column,
    )

    write_rows_to_csv(merged_rows, output_path, dialect, output_encoding)
    print(f"Output: {output_path}")

    if args.paste_to:
        paste_target = Path(args.paste_to)
        pasted_rows = paste_rows_to_csv(
            rows=merged_rows,
            target_path=paste_target,
            start_row=args.paste_start_row,
            mode=args.paste_mode,
            encoding=args.encoding,
            output_encoding=output_encoding,
            backup=not args.no_backup,
        )
        print(f"Done. Pasted {pasted_rows} row(s), {total_rows} data row(s).")
        print(f"Paste target: {paste_target}")
        if not args.no_backup:
            print(f"Backup: {paste_target.with_suffix(paste_target.suffix + '.bak')}")
    elif args.paste_to_excel:
        paste_target = Path(args.paste_to_excel)
        start_col = parse_excel_column(args.paste_start_col)
        pasted_rows = paste_rows_to_excel(
            rows=merged_rows,
            target_path=paste_target,
            sheet_name=args.excel_sheet,
            start_row=args.paste_start_row,
            start_col=start_col,
            mode=args.paste_mode,
        )
        sheet_display = args.excel_sheet or "(active sheet)"
        print(f"Done. Pasted {pasted_rows} row(s), {total_rows} data row(s) to Excel.")
        print(f"Excel target: {paste_target}")
        print(f"Sheet       : {sheet_display}")
    else:
        print(f"Done. Merged {len(files)} file(s), {total_rows} data row(s).")


if __name__ == "__main__":
    main()
