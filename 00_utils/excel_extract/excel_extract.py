"""
excel_extract.py - Excelファイルから指定シートのデータ行を抽出・マージするツール

【概要】
指定フォルダ内のExcelファイルを対象に、指定したシート・開始行から
特定列が連続して空になるまでデータを読み込み、全ファイルをマージしてCSVに出力する。
設計書に記載されたファイル名一覧などを抽出する用途を想定している。

【使い方】
  python excel_extract.py <フォルダパス> --sheet <シート名またはパターン> --stop-col <空判定列名> [オプション]

【引数】
  folder               対象フォルダのパス
  --sheet, -s          読み込むシート名（必須）。ワイルドカード使用可（例: ファイル*、*一覧）
  --stop-col, -c       この列が連続して空になったら読み込みを終了する列名（必須）
  --header-row         ヘッダー行の行番号（省略時: 1）
  --start-row          データ読み込み開始行の行番号（省略時: ヘッダー行 + 1）
  --max-empty          連続空行の閾値。この行数連続して空なら終了（省略時: 10）
  --output, -o         出力CSVのパス（省略時: excel_extract_result.csv）
  --no-recursive       サブフォルダを含めない

【出力CSVの列】
  ファイル名（先頭列）, シート名, Excelのヘッダー行の各列名

【ワイルドカード仕様】
  *  : 任意の文字列にマッチ（例: ファイル* → ファイル一覧、ファイル情報 など）
  ?  : 任意の1文字にマッチ（例: シート? → シートA、シートB など）
  複数シートがマッチした場合はすべて処理し、出力の「シート名」列で区別する。

【連続空行の仕様】
  空判定列が空の行は出力に含めず、連続カウントを加算する。
  非空行が現れたらカウントをリセットしてデータ取得を継続する。
  連続空行が --max-empty に達した時点で読み込みを終了する。

【依存ライブラリ】
  pip install openpyxl

【使用例】
  # ワイルドカードでシート名を指定
  python excel_extract.py C:/work/設計書 --sheet ファイル* --stop-col ファイル名

  # 連続10行空で終了（デフォルト）、出力先を指定
  python excel_extract.py C:/work/設計書 --sheet *一覧 --stop-col ファイル名 --output result.csv

  # ヘッダーが2行目、データが3行目からの場合
  python excel_extract.py C:/work/設計書 --sheet ファイル一覧 --stop-col ファイル名 --header-row 2 --start-row 3

  # 連続空行の閾値を5行に変更
  python excel_extract.py C:/work/設計書 --sheet ファイル* --stop-col ファイル名 --max-empty 5
"""

import argparse
import csv
import fnmatch
import os
import sys

try:
    import openpyxl
except ImportError:
    print("[ERROR] openpyxl がインストールされていません。: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


# -----------------------------------------------------------------------
# Excelファイル収集
# -----------------------------------------------------------------------

def collect_excel_files(folder: str, recursive: bool) -> list[str]:
    """対象フォルダからExcelファイルを収集する。"""
    files = []
    if recursive:
        for dirpath, _, filenames in os.walk(folder):
            for name in sorted(filenames):
                if name.lower().endswith(".xlsx") and not name.startswith("~$"):
                    files.append(os.path.join(dirpath, name))
    else:
        for entry in sorted(os.scandir(folder), key=lambda e: e.name):
            if entry.is_file() and entry.name.lower().endswith(".xlsx") and not entry.name.startswith("~$"):
                files.append(entry.path)
    return files


# -----------------------------------------------------------------------
# シート名マッチング
# -----------------------------------------------------------------------

def match_sheets(sheet_names: list[str], pattern: str) -> list[str]:
    """
    ワイルドカードパターンにマッチするシート名を返す。
    ワイルドカード文字（* ? [ ]）が含まれない場合は完全一致で検索する。
    """
    is_wildcard = any(c in pattern for c in ("*", "?", "[", "]"))
    if is_wildcard:
        return [name for name in sheet_names if fnmatch.fnmatch(name, pattern)]
    else:
        return [name for name in sheet_names if name == pattern]


# -----------------------------------------------------------------------
# 抽出処理
# -----------------------------------------------------------------------

def extract_from_sheet(
    ws,
    file_name: str,
    sheet_name: str,
    header_row: int,
    start_row: int,
    stop_col: str,
    max_empty: int,
) -> tuple[list[str], list[dict]]:
    """
    1シートからデータを抽出する。

    Returns:
        (headers, rows)
        headers: 列名のリスト（ヘッダー行から取得）
        rows   : 抽出データのリスト（各要素はdict: {列名: 値}）
    """
    # ヘッダー行を取得
    headers = []
    for cell in ws[header_row]:
        val = cell.value
        headers.append(str(val).strip() if val is not None else "")

    if not any(headers):
        print(f"    [WARN] ヘッダー行({header_row}行目)が空です: {file_name} / {sheet_name}", file=sys.stderr)
        return [], []

    # stop_col のインデックスを確認
    if stop_col not in headers:
        print(f"    [WARN] 空判定列 '{stop_col}' がヘッダーに見つかりません: {file_name} / {sheet_name}", file=sys.stderr)
        return [], []

    stop_col_idx = headers.index(stop_col)

    # データ行を読み込む
    rows = []
    empty_count = 0

    for row_cells in ws.iter_rows(min_row=start_row, values_only=True):
        stop_val = row_cells[stop_col_idx] if stop_col_idx < len(row_cells) else None
        is_empty = stop_val is None or str(stop_val).strip() == ""

        if is_empty:
            empty_count += 1
            if empty_count >= max_empty:
                # 連続空行が閾値に達したので終了
                break
            # 閾値未満の空行はスキップして継続
            continue

        # 非空行：カウントリセットしてデータ取得
        empty_count = 0
        row_dict = {}
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            val = row_cells[col_idx] if col_idx < len(row_cells) else None
            row_dict[header] = str(val).strip() if val is not None else ""

        rows.append(row_dict)

    return headers, rows


def extract_from_excel(
    file_path: str,
    sheet_pattern: str,
    header_row: int,
    start_row: int,
    stop_col: str,
    max_empty: int,
) -> tuple[list[str], list[dict]]:
    """
    1つのExcelファイルからパターンにマッチする全シートのデータを抽出する。

    Returns:
        (headers, rows)
        headers: 列名のリスト
        rows   : 抽出データのリスト（ファイル名・シート名付き）
    """
    file_name = os.path.basename(file_path)

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"  [WARN] ファイルを開けませんでした: {file_name} ({e})", file=sys.stderr)
        return [], []

    matched = match_sheets(wb.sheetnames, sheet_pattern)
    if not matched:
        print(f"  [WARN] パターン '{sheet_pattern}' にマッチするシートがありません: {file_name}", file=sys.stderr)
        wb.close()
        return [], []

    all_headers: list[str] = []
    all_rows: list[dict] = []

    for sheet_name in matched:
        print(f"    シート: {sheet_name}")
        ws = wb[sheet_name]
        headers, rows = extract_from_sheet(
            ws=ws,
            file_name=file_name,
            sheet_name=sheet_name,
            header_row=header_row,
            start_row=start_row,
            stop_col=stop_col,
            max_empty=max_empty,
        )

        if not headers:
            continue

        if not all_headers:
            all_headers = headers

        for row in rows:
            row["ファイル名"] = file_name
            row["シート名"] = sheet_name
            all_rows.append(row)

        print(f"      → {len(rows)} 行抽出")

    wb.close()
    return all_headers, all_rows


# -----------------------------------------------------------------------
# CSV出力
# -----------------------------------------------------------------------

def write_csv(
    all_headers: list[str],
    all_rows: list[dict],
    output_path: str,
) -> None:
    """マージ済みデータをCSVに書き出す。先頭列はファイル名・シート名。"""
    fieldnames = ["ファイル名", "シート名"] + all_headers
    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(all_rows)


# -----------------------------------------------------------------------
# メイン
# -----------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Excelファイルから指定シートのデータ行を抽出・マージしてCSV出力する"
    )
    parser.add_argument("folder", help="対象フォルダのパス")
    parser.add_argument("--sheet", "-s", required=True, help="読み込むシート名（ワイルドカード使用可: 例 ファイル*）")
    parser.add_argument("--stop-col", "-c", required=True, help="この列が連続して空になったら読み込みを終了する列名")
    parser.add_argument("--header-row", type=int, default=1, help="ヘッダー行の行番号（デフォルト: 1）")
    parser.add_argument("--start-row", type=int, default=None, help="データ読み込み開始行の行番号（省略時: ヘッダー行 + 1）")
    parser.add_argument("--max-empty", type=int, default=10, help="終了と判断する連続空行数（デフォルト: 10）")
    parser.add_argument("--output", "-o", default="excel_extract_result.csv", help="出力CSVのパス（デフォルト: excel_extract_result.csv）")
    parser.add_argument("--no-recursive", action="store_true", help="サブフォルダを含めない")

    args = parser.parse_args()

    # 開始行の補完
    start_row = args.start_row if args.start_row is not None else args.header_row + 1

    # 開始行とヘッダー行の整合チェック
    if start_row <= args.header_row:
        print(f"[ERROR] --start-row ({start_row}) は --header-row ({args.header_row}) より大きい値にしてください。", file=sys.stderr)
        sys.exit(1)

    # フォルダ確認
    if not os.path.isdir(args.folder):
        print(f"[ERROR] フォルダが見つかりません: {args.folder}", file=sys.stderr)
        sys.exit(1)

    recursive = not args.no_recursive
    excel_files = collect_excel_files(args.folder, recursive)

    if not excel_files:
        print("対象のExcelファイルが見つかりませんでした。")
        sys.exit(0)

    print(f"対象フォルダ  : {args.folder}")
    print(f"対象ファイル数: {len(excel_files)} 件")
    print(f"シートパターン: {args.sheet}")
    print(f"ヘッダー行    : {args.header_row} 行目")
    print(f"データ開始行  : {start_row} 行目")
    print(f"空判定列      : {args.stop_col}")
    print(f"連続空行閾値  : {args.max_empty} 行")
    print()

    # 全ファイルから抽出
    all_headers: list[str] = []
    all_rows: list[dict] = []
    total_rows = 0
    skipped_files = 0

    for file_path in excel_files:
        file_name = os.path.basename(file_path)
        print(f"  抽出中: {file_name}")

        headers, rows = extract_from_excel(
            file_path=file_path,
            sheet_pattern=args.sheet,
            header_row=args.header_row,
            start_row=start_row,
            stop_col=args.stop_col,
            max_empty=args.max_empty,
        )

        if not headers:
            skipped_files += 1
            continue

        if not all_headers:
            all_headers = headers

        all_rows.extend(rows)
        total_rows += len(rows)

    print()

    if not all_rows:
        print("抽出できたデータがありませんでした。")
        sys.exit(0)

    write_csv(all_headers, all_rows, args.output)

    print(f"完了: 合計 {total_rows} 行を抽出しました。（スキップ: {skipped_files} ファイル）")
    print(f"出力先: {args.output}")


if __name__ == "__main__":
    main()
