"""
keyword_search.py - キーワード検索ツール

【概要】
指定したフォルダ配下のファイルを対象に、特定のキーワードを含む行・セルを検索し、
ファイル名と該当箇所の内容を一覧で出力する。
テキスト系ファイル（.txt/.csv/.md 等）、Excel（.xlsx/.xls）、Word（.docx）に対応。
.xlsx は openpyxl、.xls は xlrd を使用する（pip install openpyxl xlrd）。

【使い方】
  単一フォルダ指定:
    python keyword_search.py <フォルダパス> <キーワード> [オプション]

  複数フォルダ指定（テキストファイル）:
    python keyword_search.py <キーワード> --folders-file <テキストファイルパス> [オプション]

  ※ テキストファイルは1行1フォルダパスで記載（空行・#始まり行はスキップ）

【オプション】
  --folders-file テキストファイルから対象フォルダを複数読み込む（1行1フォルダ）
                 行末に |no-recursive または |recursive を付けるとフォルダごとに個別指定可能
                 例: C:\work\設計書\PJA|no-recursive
  --ext          対象拡張子をカンマ区切りで指定（省略時: .txt,.csv,.md,.xlsx,.xls,.docx）
                 例: --ext .xlsx,.csv
  --filename     ファイル名フィルタ（ワイルドカード対応、省略時: 全ファイル）
                 例: --filename "IF*.xlsx"
  --sheet        Excelのシート名フィルタ（ワイルドカード対応、省略時: 全シート）
                 例: --sheet "一覧*"
  --no-recursive サブフォルダを含めない（省略時: サブフォルダも含める）
  --ignore-case  大文字/小文字を区別しない（省略時: 区別する）
  --outdir       出力先フォルダのパス（省略時: このスクリプトと同じフォルダ）
  --encoding     CSVの文字コード（省略時: cp932＝Excelでそのまま開いて文字化けしない）

【出力列】
  ファイルパス, ファイル名, 場所, 内容

  「場所」の表記：
    テキスト/CSV系 → 行番号（例: L12）
    Excel          → シート名とセルアドレス（例: Sheet1:B5）
    Word           → 段落番号（例: P3）

【出力ファイル名】
  keyword_search_YYYYMMDD_HHMMSS.csv（タイムスタンプ付き）

【使用例】
  # 単一フォルダ
  python keyword_search.py C:/work/設計書 "ユーザーID"
  python keyword_search.py C:/work/設計書 "ユーザーID" --ext .xlsx,.docx
  python keyword_search.py C:/work/設計書 "エラー" --ignore-case --outdir C:/work/output

  # Excelのシートを絞り込み
  python keyword_search.py C:/work/設計書 "ユーザーID" --ext .xlsx --sheet "一覧*"
  python keyword_search.py C:/work/設計書 "エラー" --ext .xlsx --sheet "Sheet1"

  # 複数フォルダ（テキストファイル指定）
  python keyword_search.py "ユーザーID" --folders-file folders.txt
  python keyword_search.py "エラー" --folders-file folders.txt --ext .xlsx --ignore-case
"""

import argparse
import csv
import fnmatch
import os
import sys
import unicodedata
from datetime import datetime


# ---------------------------------------------------------------------------
# 検索ロジック（ファイル種別ごと）
# ---------------------------------------------------------------------------

def search_text_file(filepath: str, keyword: str, ignore_case: bool) -> list[dict]:
    """テキスト系ファイル（.txt / .csv / .md 等）を行単位で検索する。"""
    results = []
    # 文字コードを順に試す（BOMあり UTF-8 → UTF-8 → cp932）
    for enc in ("utf-8-sig", "utf-8", "cp932"):
        try:
            with open(filepath, encoding=enc, errors="strict") as f:
                lines = f.readlines()
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:
        # すべてのエンコードで失敗した場合はスキップ
        return results

    needle = keyword.lower() if ignore_case else keyword
    for lineno, line in enumerate(lines, start=1):
        haystack = line.lower() if ignore_case else line
        if needle in haystack:
            results.append({
                "場所": f"L{lineno}",
                "内容": line.rstrip("\n"),
            })
    return results


def _col_letter(col_idx: int) -> str:
    """0始まりの列インデックスをExcelの列名（A, B, ..., Z, AA, ...）に変換する。"""
    letter = ""
    col_idx += 1  # 1始まりに変換
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        letter = chr(ord("A") + remainder) + letter
    return letter


def search_excel_file(filepath: str, keyword: str, ignore_case: bool,
                      sheet_pattern: str | None = None) -> list[dict]:
    """Excel ファイル（.xlsx / .xls）をセル単位で検索する。

    .xlsx は openpyxl、.xls は xlrd で処理する。

    Args:
        sheet_pattern: シート名フィルタ（ワイルドカード対応）。None の場合は全シートを対象。
    """
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".xls":
        return _search_xls_file(filepath, keyword, ignore_case, sheet_pattern)
    return _search_xlsx_file(filepath, keyword, ignore_case, sheet_pattern)


def _search_xlsx_file(filepath: str, keyword: str, ignore_case: bool,
                      sheet_pattern: str | None) -> list[dict]:
    """openpyxl で .xlsx を検索する。"""
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        print("[WARN] openpyxl が未インストールのため Excel ファイルをスキップします。"
              "  pip install openpyxl", file=sys.stderr)
        return []

    results = []
    needle = keyword.lower() if ignore_case else keyword

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        print(f"[WARN] Excel を開けませんでした ({os.path.basename(filepath)}): {e}",
              file=sys.stderr)
        return results

    try:
        for sheet_name in wb.sheetnames:
            if sheet_pattern and not fnmatch.fnmatch(sheet_name, sheet_pattern):
                continue
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    cell_str = str(cell.value)
                    haystack = cell_str.lower() if ignore_case else cell_str
                    if needle in haystack:
                        results.append({
                            "場所": f"{sheet_name}:{cell.coordinate}",
                            "内容": cell_str,
                        })
    finally:
        wb.close()
    return results


def _search_xls_file(filepath: str, keyword: str, ignore_case: bool,
                     sheet_pattern: str | None) -> list[dict]:
    """xlrd で .xls を検索する。"""
    try:
        import xlrd  # noqa: PLC0415
    except ImportError:
        print("[WARN] xlrd が未インストールのため .xls ファイルをスキップします。"
              "  pip install xlrd", file=sys.stderr)
        return []

    results = []
    needle = keyword.lower() if ignore_case else keyword

    try:
        wb = xlrd.open_workbook(filepath)
    except Exception as e:  # noqa: BLE001
        print(f"[WARN] Excel(.xls) を開けませんでした ({os.path.basename(filepath)}): {e}",
              file=sys.stderr)
        return results

    for sheet_name in wb.sheet_names():
        if sheet_pattern and not fnmatch.fnmatch(sheet_name, sheet_pattern):
            continue
        ws = wb.sheet_by_name(sheet_name)
        for row_idx in range(ws.nrows):
            for col_idx in range(ws.ncols):
                cell_value = ws.cell_value(row_idx, col_idx)
                if cell_value == "" or cell_value is None:
                    continue
                cell_str = str(cell_value)
                haystack = cell_str.lower() if ignore_case else cell_str
                if needle in haystack:
                    coordinate = f"{_col_letter(col_idx)}{row_idx + 1}"
                    results.append({
                        "場所": f"{sheet_name}:{coordinate}",
                        "内容": cell_str,
                    })
    return results


def search_word_file(filepath: str, keyword: str, ignore_case: bool) -> list[dict]:
    """Word ファイル（.docx）を段落単位で検索する。"""
    try:
        from docx import Document  # noqa: PLC0415
    except ImportError:
        print("[WARN] python-docx が未インストールのため Word ファイルをスキップします。"
              "  pip install python-docx", file=sys.stderr)
        return []

    results = []
    needle = keyword.lower() if ignore_case else keyword

    try:
        doc = Document(filepath)
    except Exception as e:  # noqa: BLE001
        print(f"[WARN] Word を開けませんでした ({os.path.basename(filepath)}): {e}",
              file=sys.stderr)
        return results

    for para_no, para in enumerate(doc.paragraphs, start=1):
        text = para.text
        if not text.strip():
            continue
        haystack = text.lower() if ignore_case else text
        if needle in haystack:
            results.append({
                "場所": f"P{para_no}",
                "内容": text,
            })
    return results


# ---------------------------------------------------------------------------
# ファイル収集
# ---------------------------------------------------------------------------

DEFAULT_EXTENSIONS = {".txt", ".csv", ".md", ".xlsx", ".xls", ".docx"}


def collect_files(root_dir: str, extensions: set[str], filename_pattern: str | None,
                  recursive: bool) -> list[str]:
    """対象ファイルのパス一覧を返す。"""
    filepaths = []

    if recursive:
        walker = os.walk(root_dir)
    else:
        entries = os.scandir(root_dir)
        walker = [(root_dir, [], [e.name for e in entries if e.is_file()])]

    for dirpath, _, filenames in walker:
        for filename in sorted(filenames):
            ext = os.path.splitext(filename)[1].lower()
            if ext not in extensions:
                continue
            if filename_pattern and not fnmatch.fnmatch(filename, filename_pattern):
                continue
            filepaths.append(os.path.join(dirpath, filename))

    return filepaths


# ---------------------------------------------------------------------------
# 検索ディスパッチ
# ---------------------------------------------------------------------------

TEXT_EXTENSIONS = {".txt", ".csv", ".md", ".log", ".ini", ".yaml", ".yml", ".json"}
EXCEL_EXTENSIONS = {".xlsx", ".xls"}
WORD_EXTENSIONS = {".docx"}


def search_file(filepath: str, keyword: str, ignore_case: bool,
                sheet_pattern: str | None = None) -> list[dict]:
    """拡張子に応じて適切な検索関数を呼び出す。"""
    ext = os.path.splitext(filepath)[1].lower()
    if ext in EXCEL_EXTENSIONS:
        return search_excel_file(filepath, keyword, ignore_case, sheet_pattern)
    if ext in WORD_EXTENSIONS:
        return search_word_file(filepath, keyword, ignore_case)
    return search_text_file(filepath, keyword, ignore_case)


# ---------------------------------------------------------------------------
# 出力
# ---------------------------------------------------------------------------

def _display_width(text: str) -> int:
    """全角文字を幅2、半角を幅1として表示幅を返す。"""
    width = 0
    for ch in text:
        eaw = unicodedata.east_asian_width(ch)
        width += 2 if eaw in ("W", "F") else 1
    return width


def _ljust_wide(text: str, width: int) -> str:
    """全角文字を考慮して左詰めパディングした文字列を返す。"""
    pad = width - _display_width(text)
    return text + " " * max(pad, 0)


def print_results(records: list[dict]) -> None:
    """コンソールに結果を出力する。"""
    if not records:
        return
    # 全角文字を考慮した表示幅で列幅を計算
    max_file = max(_display_width(r["ファイル名"]) for r in records)
    max_loc  = max(_display_width(r["場所"])       for r in records)
    header = f"{_ljust_wide('ファイル名', max_file)}  {_ljust_wide('場所', max_loc)}  内容"
    print(header)
    print("-" * min(len(header) + 20, 120))
    for r in records:
        print(f"{_ljust_wide(r['ファイル名'], max_file)}  {_ljust_wide(r['場所'], max_loc)}  {r['内容']}")


def write_csv(records: list[dict], output_path: str, encoding: str = "cp932") -> None:
    """結果をCSVに書き出す。"""
    fieldnames = ["ファイルパス", "ファイル名", "場所", "内容"]
    with open(output_path, "w", newline="", encoding=encoding) as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)


# ---------------------------------------------------------------------------
# メイン
# ---------------------------------------------------------------------------

def load_folders_file(filepath: str) -> list[tuple[str, bool | None]]:
    """フォルダリストファイルを読み込む。空行・#始まり行はスキップ。

    各行の形式:
      <フォルダパス>                  → recursive はグローバル設定に従う（None）
      <フォルダパス>|no-recursive     → このフォルダはサブフォルダを含めない
      <フォルダパス>|recursive        → このフォルダはサブフォルダを含める

    Returns:
        (フォルダパス, recursive設定) のリスト。recursive設定は True/False/None。
        None はグローバル設定に従うことを意味する。
    """
    folders = []
    for enc in ("utf-8-sig", "utf-8", "cp932"):
        try:
            with open(filepath, encoding=enc, errors="strict") as f:
                lines = f.readlines()
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:
        print(f"[ERROR] フォルダリストファイルを読み込めませんでした: {filepath}", file=sys.stderr)
        sys.exit(1)

    for line in lines:
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if "|" in line:
            path, flag = line.rsplit("|", maxsplit=1)
            path = path.strip()
            flag = flag.strip().lower()
            if flag == "no-recursive":
                folders.append((path, False))
            elif flag == "recursive":
                folders.append((path, True))
            else:
                # 不明なフラグは無視してグローバル設定に従う
                print(f"[WARN] 不明なフラグを無視します（使用可能: no-recursive / recursive）: |{flag}",
                      file=sys.stderr)
                folders.append((path, None))
        else:
            folders.append((line, None))
    return folders


def main() -> None:
    parser = argparse.ArgumentParser(
        description="フォルダ配下のファイルからキーワードを検索し、該当行・セルを一覧表示する"
    )
    parser.add_argument("folder", nargs="?", default="",
                        help="対象フォルダのパス（--folders-file 指定時は省略可）")
    parser.add_argument("keyword", help="検索するキーワード")
    parser.add_argument("--folders-file", default="",
                        help="対象フォルダを列挙したテキストファイルのパス（1行1フォルダ）")
    parser.add_argument("--ext", default="",
                        help="対象拡張子をカンマ区切りで指定（例: .xlsx,.csv）。省略時はデフォルト拡張子を使用")
    parser.add_argument("--filename", default="",
                        help="ファイル名フィルタ（ワイルドカード対応、例: IF*.xlsx）")
    parser.add_argument("--sheet", default="",
                        help="Excelのシート名フィルタ（ワイルドカード対応、例: 一覧*）。省略時は全シートを対象")
    parser.add_argument("--no-recursive", action="store_true",
                        help="サブフォルダを含めない")
    parser.add_argument("--ignore-case", action="store_true",
                        help="大文字/小文字を区別しない")
    parser.add_argument("--outdir", default="",
                        help="出力先フォルダのパス（省略時: このスクリプトと同じフォルダ）")
    parser.add_argument("--encoding", default="cp932",
                        help="CSVの文字コード（デフォルト: cp932）")

    args = parser.parse_args()

    # 対象フォルダリストの決定
    # target_folders は (フォルダパス, recursive設定) のリスト
    # recursive設定が None の場合はグローバルの --no-recursive に従う
    if args.folders_file:
        # テキストファイルから複数フォルダを読み込む
        if not os.path.isfile(args.folders_file):
            print(f"[ERROR] フォルダリストファイルが見つかりません: {args.folders_file}", file=sys.stderr)
            sys.exit(1)
        target_folders = load_folders_file(args.folders_file)
        if not target_folders:
            print(f"[ERROR] フォルダリストファイルに有効なフォルダが1件もありません: {args.folders_file}",
                  file=sys.stderr)
            sys.exit(1)
    elif args.folder:
        # 単一フォルダ指定（recursive設定はグローバルに従う）
        target_folders = [(args.folder, None)]
    else:
        print("[ERROR] 対象フォルダを指定してください（引数 or --folders-file）", file=sys.stderr)
        parser.print_help()
        sys.exit(1)

    # フォルダの存在確認
    invalid = [f for f, _ in target_folders if not os.path.isdir(f)]
    if invalid:
        for f in invalid:
            print(f"[ERROR] フォルダが見つかりません: {f}", file=sys.stderr)
        sys.exit(1)

    # 拡張子セットの構築
    if args.ext:
        extensions = {e.strip().lower() for e in args.ext.split(",") if e.strip()}
    else:
        extensions = DEFAULT_EXTENSIONS

    global_recursive = not args.no_recursive
    filename_pattern = args.filename if args.filename else None
    sheet_pattern = args.sheet if args.sheet else None

    # 出力先フォルダの決定（省略時: スクリプトと同じフォルダ）
    outdir = args.outdir if args.outdir else os.path.dirname(os.path.abspath(__file__))
    if not os.path.isdir(outdir):
        print(f"[ERROR] 出力先フォルダが見つかりません: {outdir}", file=sys.stderr)
        sys.exit(1)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(outdir, f"keyword_search_{timestamp}.csv")

    print(f"対象フォルダ   : {len(target_folders)} 件")
    for f, rec in target_folders:
        effective = global_recursive if rec is None else rec
        label = "（個別指定）" if rec is not None else ""
        print(f"  {f}  [サブフォルダ: {'含める' if effective else '含めない'}{label}]")
    print(f"キーワード     : {args.keyword}")
    print(f"拡張子フィルタ : {sorted(extensions)}")
    print(f"ファイル名     : {filename_pattern if filename_pattern else '（指定なし）'}")
    print(f"シート名       : {sheet_pattern if sheet_pattern else '（指定なし＝全シート）'}")
    print(f"大文字小文字   : {'区別しない' if args.ignore_case else '区別する'}")
    print(f"出力先         : {output_path}")
    print()

    # ファイル収集（フォルダごとに recursive 設定を適用）
    all_filepaths: list[str] = []
    for folder, rec in target_folders:
        recursive = global_recursive if rec is None else rec
        all_filepaths.extend(collect_files(folder, extensions, filename_pattern, recursive))

    if not all_filepaths:
        print("対象ファイルが見つかりませんでした。")
        sys.exit(0)

    # 検索
    all_records: list[dict] = []
    for filepath in all_filepaths:
        hits = search_file(filepath, args.keyword, args.ignore_case, sheet_pattern)
        dirpath = os.path.dirname(filepath)
        filename = os.path.basename(filepath)
        for hit in hits:
            all_records.append({
                "ファイルパス": dirpath,
                "ファイル名": filename,
                "場所": hit["場所"],
                "内容": hit["内容"],
            })

    # 結果出力
    print(f"検索ファイル数 : {len(all_filepaths)}")
    print(f"ヒット件数     : {len(all_records)}")
    print()

    if all_records:
        print_results(all_records)
    else:
        print(f"「{args.keyword}」に該当する箇所は見つかりませんでした。")

    # CSV出力（常に実行）
    write_csv(all_records, output_path, encoding=args.encoding)
    print()
    print(f"出力先: {output_path}")


if __name__ == "__main__":
    main()
